"""Unit tests for WorkbookAnalyzer."""

import os
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
import pytest
from excel_copilot.analyzer.workbook_analyzer import WorkbookAnalyzer


@pytest.fixture
def sample_excel_file(tmp_path):
    file_path = str(tmp_path / "sample_sales.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"

    # Add header and data
    headers = ["Date", "Revenue", "Cost", "Profit"]
    ws.append(headers)

    ws.append(["2026-01-01", 1000, 600, "=B2-C2"])
    ws.append(["2026-01-02", 1500, 800, "=B3-C3"])
    ws.append(["2026-01-03", 2000, 1100, "=B4-C4"])

    # Create Native Excel Table
    tab = Table(displayName="SalesData", ref="A1:D4")
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True)
    tab.tableStyleInfo = style
    ws.add_table(tab)

    wb.save(file_path)
    wb.close()
    return file_path


def test_workbook_analyzer(sample_excel_file):
    model = WorkbookAnalyzer.analyze(sample_excel_file)

    assert model.filename == "sample_sales.xlsx"
    assert len(model.worksheets) == 1

    ws = model.worksheets[0]
    assert ws.name == "Sales"
    assert len(ws.tables) == 1

    tbl = ws.tables[0]
    assert tbl.name == "SalesData"
    assert len(tbl.columns) == 4
    assert [c.name for c in tbl.columns] == ["Date", "Revenue", "Cost", "Profit"]


def test_generate_system_prompt(sample_excel_file):
    model = WorkbookAnalyzer.analyze(sample_excel_file)
    prompt = WorkbookAnalyzer.generate_system_prompt(model)

    assert "SYSTEM INSTRUCTIONS: EXCEL AI COPILOT" in prompt
    assert "SalesData" in prompt
    assert "excel-action" in prompt
