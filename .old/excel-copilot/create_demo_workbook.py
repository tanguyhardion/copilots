"""Generates sample Financial_Dashboard.xlsx workbook for testing."""

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

def create_demo():
    wb = openpyxl.Workbook()

    # Sheet 1: Sales
    ws_sales = wb.active
    ws_sales.title = "Sales"
    ws_sales.append(["Date", "Region", "Revenue", "Cost"])
    ws_sales.append(["2026-01-01", "North", 15000, 9000])
    ws_sales.append(["2026-01-02", "South", 22000, 14000])
    ws_sales.append(["2026-01-03", "East", 18000, 11000])
    ws_sales.append(["2026-01-04", "West", 25000, 15000])

    tab_sales = Table(displayName="SalesTable", ref="A1:D5")
    ws_sales.add_table(tab_sales)

    # Sheet 2: Dashboard
    ws_dash = wb.create_sheet(title="Dashboard")
    ws_dash.append(["KPI Metric", "Target", "Actual"])
    ws_dash.append(["Total Revenue", 100000, 80000])
    ws_dash.append(["Gross Margin", 0.40, 0.38])

    tab_dash = Table(displayName="KPIDashboard", ref="A1:C3")
    ws_dash.add_table(tab_dash)

    wb.save("Financial_Dashboard.xlsx")
    wb.close()
    print("Created Financial_Dashboard.xlsx successfully!")

if __name__ == "__main__":
    create_demo()
