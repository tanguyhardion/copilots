import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from copilots_app.services.cv.generator import fmt_cert_expiry, fmt_date, generate_cv
from copilots_app.services.cv.i18n import get_cv_strings
from copilots_app.services.cv.pptx_generator import generate_pptx_cv


SAMPLE_CV = {
    "personal_info": {
        "first_name": "Jean",
        "last_name": "Dupont",
        "phone": "0601020304",
        "email": "jean.dupont@example.com",
        "years_experience_bucket": "4_to_9",
    },
    "tender_info": {"proposed_role": "Manager"},
    "profile": "Delivery lead with public-sector experience.",
    "work_experience": [
        {
            "date_from": "2020-01",
            "date_to": "Present",
            "job_title": "Manager",
            "organisation": "Contoso",
        }
    ],
    "project_experience": [
        {
            "date_from": "2021-01",
            "date_to": "Present",
            "role": "Lead",
            "client": "Client",
            "project_title": "Programme",
            "responsibilities": "Delivered outcomes.",
        }
    ],
    "education": [
        {
            "date_from": "2015-09",
            "date_to": "2017-06",
            "qualification_title": "Master",
        }
    ],
    "experience_overview": [
        {
            "date_from": "2021-01",
            "date_to": "Present",
            "role": "Lead",
            "responsibilities": "Delivered outcomes.",
        }
    ],
    "personal_skills": {
        "communication": ["Communication"],
        "organisational_managerial": ["Planning"],
        "computer_skills": ["Python"],
        "certifications": [{"year": 2024, "title": "PMP"}],
    },
    "languages": {"mother_tongue": ["French"], "other": []},
}


class FakeRun:
    def __init__(self):
        self.font = SimpleNamespace(
            name=None,
            size=None,
            bold=None,
            color=SimpleNamespace(rgb=None),
        )


class FakeParagraph:
    def __init__(self, text=""):
        self._text = ""
        self.runs = []
        self.space_after = None
        self.space_before = None
        self.text = text

    @property
    def text(self):
        return self._text

    @text.setter
    def text(self, value):
        self._text = value
        self.runs = [FakeRun()] if value else []


class FakeTextFrame:
    def __init__(self, paragraph_count):
        self.paragraphs = [FakeParagraph() for _ in range(paragraph_count)]

    def add_paragraph(self):
        paragraph = FakeParagraph()
        self.paragraphs.append(paragraph)
        return paragraph


class FakeCell:
    def __init__(self):
        self.text_frame = FakeTextFrame(1)

    @property
    def text(self):
        return self.text_frame.paragraphs[0].text

    @text.setter
    def text(self, value):
        self.text_frame = FakeTextFrame(1)
        self.text_frame.paragraphs[0].text = value


class FakeTable:
    def __init__(self, rows, cols):
        self.rows = [object() for _ in range(rows)]
        self.columns = [object() for _ in range(cols)]
        self._cells = [[FakeCell() for _ in range(cols)] for _ in range(rows)]

    def cell(self, row, col):
        return self._cells[row][col]


class FakeShape:
    def __init__(self, name, paragraph_count=None, table=None):
        self.name = name
        self.has_table = table is not None
        if table is not None:
            self.table = table
        else:
            self.text_frame = FakeTextFrame(paragraph_count or 1)


class FakePresentation:
    def __init__(self):
        self.header_shape = FakeShape("Text Placeholder 2", paragraph_count=5)
        self.profile_shape = FakeShape("TextBox 7", paragraph_count=1)
        self.skills_table = FakeTable(2, 3)
        self.experience_table = FakeTable(5, 2)
        self.slides = [
            SimpleNamespace(
                shapes=[
                    self.header_shape,
                    self.profile_shape,
                    FakeShape("Table 72", table=self.skills_table),
                    FakeShape("Table 71", table=self.experience_table),
                ]
            )
        ]
        self.saved_path = None

    def save(self, output_path):
        self.saved_path = output_path


class CVGeneratorTests(unittest.TestCase):
    def test_generate_cv_respects_section_toggles_and_language(self):
        fake_doc = Mock()
        with (
            patch("copilots_app.services.cv.generator.setup_document", return_value=fake_doc),
            patch("copilots_app.services.cv.generator.build_header") as build_header,
            patch("copilots_app.services.cv.generator.build_footer") as build_footer,
            patch("copilots_app.services.cv.generator.build_personal_info") as build_personal_info,
            patch("copilots_app.services.cv.generator.build_proposed_role") as build_proposed_role,
            patch("copilots_app.services.cv.generator.build_professional_experience") as build_professional_experience,
            patch("copilots_app.services.cv.generator.build_work_experience") as build_work_experience,
            patch("copilots_app.services.cv.generator.build_profile") as build_profile,
            patch("copilots_app.services.cv.generator.build_requirements_matrix") as build_requirements_matrix,
            patch("copilots_app.services.cv.generator.build_experience_overview") as build_experience_overview,
            patch("copilots_app.services.cv.generator.build_project_experience") as build_project_experience,
            patch("copilots_app.services.cv.generator.build_education") as build_education,
            patch("copilots_app.services.cv.generator.build_personal_skills") as build_personal_skills,
        ):
            generate_cv(
                SAMPLE_CV,
                "/tmp/sample.docx",
                language="fr",
                sections={"profile": False, "education": False},
            )

        build_header.assert_called_once()
        build_footer.assert_called_once()
        build_personal_info.assert_called_once()
        build_proposed_role.assert_called_once()
        build_professional_experience.assert_called_once()
        build_work_experience.assert_called_once()
        build_requirements_matrix.assert_called_once()
        build_experience_overview.assert_called_once()
        build_project_experience.assert_called_once()
        build_personal_skills.assert_called_once()
        build_profile.assert_not_called()
        build_education.assert_not_called()
        self.assertEqual(build_header.call_args.args[3]["present"], "Présent")
        fake_doc.save.assert_called_once_with("/tmp/sample.docx")

    def test_docx_format_helpers_localize_output(self):
        strings = get_cv_strings("fr")
        self.assertEqual(fmt_date("2021-01", strings), "Janv. 2021")
        self.assertEqual(fmt_date("Present", strings), "Présent")
        self.assertEqual(fmt_cert_expiry({"expiry_year": None}, strings), "(perpétuelle)")
        self.assertEqual(
            strings["organisational_managerial_skills"],
            "Compétences organisationnelles\n/ managériales",
        )

    def test_skill_bullets_multiline_stays_in_left_column(self):
        from docx import Document
        from copilots_app.services.cv.generator import _skill_bullets
        doc = Document()
        _skill_bullets(
            doc,
            "Compétences organisationnelles\n/ managériales",
            ["Bullet 1", "Bullet 2"],
        )
        self.assertEqual(doc.paragraphs[0].text, "Compétences organisationnelles")
        self.assertEqual(doc.paragraphs[1].text, "/ managériales\t●  Bullet 1")
        self.assertEqual(doc.paragraphs[2].text, "●  Bullet 2")

    def test_generate_pptx_localizes_labels_and_dates(self):
        fake_prs = FakePresentation()
        with (
            patch("copilots_app.services.cv.pptx_generator._resolve_template_path", return_value="template.pptx"),
            patch("copilots_app.services.cv.pptx_generator.pptx.Presentation", return_value=fake_prs),
        ):
            with tempfile.TemporaryDirectory() as tmp:
                output_path = str(Path(tmp) / "cv.pptx")
                generate_pptx_cv(SAMPLE_CV, output_path, language="fr")

        self.assertEqual(fake_prs.saved_path, output_path)
        self.assertEqual(fake_prs.header_shape.text_frame.paragraphs[3].text, "Mobile: 0601020304")
        self.assertEqual(fake_prs.header_shape.text_frame.paragraphs[4].text, "E-mail: jean.dupont@example.com")
        self.assertEqual(
            fake_prs.experience_table.cell(0, 0).text_frame.paragraphs[0].text,
            "Janv. 2021 – Présent",
        )

    def test_generate_pptx_respects_section_toggles(self):
        fake_prs = FakePresentation()
        with (
            patch("copilots_app.services.cv.pptx_generator._resolve_template_path", return_value="template.pptx"),
            patch("copilots_app.services.cv.pptx_generator.pptx.Presentation", return_value=fake_prs),
        ):
            with tempfile.TemporaryDirectory() as tmp:
                output_path = str(Path(tmp) / "cv.pptx")
                generate_pptx_cv(
                    SAMPLE_CV,
                    output_path,
                    sections={
                        "header_contact": False,
                        "profile": False,
                        "skills": False,
                        "relevant_experience": False,
                    },
                )

        self.assertTrue(all(not paragraph.text for paragraph in fake_prs.header_shape.text_frame.paragraphs))
        self.assertEqual(fake_prs.profile_shape.text_frame.paragraphs[0].text, "")
        self.assertEqual(fake_prs.skills_table.cell(0, 0).text_frame.paragraphs[0].text, "")
        self.assertEqual(fake_prs.experience_table.cell(0, 0).text_frame.paragraphs[0].text, "")


if __name__ == "__main__":
    unittest.main()
